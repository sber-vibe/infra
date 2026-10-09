"""
Генератор синтетического пакета документов на субсидию для приёмочного тестирования стенда.

Все организации, люди и реквизиты вымышлены. Файлы имеют нейтральные имена (1.pdf, scan4.jpg),
поэтому тип можно определить только по содержимому (требование Д9).
В пакет намеренно заложены ошибки из сценария ТЗ — см. expected.json.

Зависимости (пермиссивные лицензии): pip install reportlab pillow python-docx
Запуск:   python testdata/generate.py [--out testdata/package] [--today 2026-10-09]
"""

from __future__ import annotations

import argparse
import io
import json
import random
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FONT_REGULAR = "C:/Windows/Fonts/arial.ttf"
FONT_BOLD = "C:/Windows/Fonts/arialbd.ttf"
for candidate in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",):
    if not Path(FONT_REGULAR).exists() and Path(candidate).exists():
        FONT_REGULAR = FONT_BOLD = candidate


def inn10(prefix9: str) -> str:
    """ИНН юрлица с корректной контрольной цифрой."""
    weights = [2, 4, 10, 3, 5, 9, 4, 6, 8]
    check = sum(int(d) * w for d, w in zip(prefix9, weights)) % 11 % 10
    return prefix9 + str(check)


ORG = "ООО «Ромашка-Тех»"
INN = inn10("616123456")            # 6161234567 → контрольная цифра пересчитана
INN_WRONG = INN[:-1] + str((int(INN[-1]) + 1) % 10)  # ошибка в выписке: отличается последняя цифра
OGRN = "1196196012345"
KPP = "616101001"
DIRECTOR = "Иванов Иван Иванович"
DIRECTOR_SHORT = "Иванов И. И."
ADDRESS = "344002, г. Ростов-на-Дону, ул. Большая Садовая, д. 10, оф. 5"
PROJECT = "Внедрение цифровой платформы учёта производства"
REQUESTED = 850_000
ESTIMATE_ITEMS = [
    ("Серверное оборудование", 1, 320_000),
    ("Лицензии ПО (12 мес.)", 4, 45_000),
    ("Разработка интеграций", 1, 190_000),
    ("Обучение персонала", 10, 11_000),
]
ESTIMATE_TOTAL = sum(q * p for _, q, p in ESTIMATE_ITEMS)  # 800 000 — меньше запрошенной суммы


def money(v: int) -> str:
    return f"{v:,}".replace(",", " ") + " руб."


# ---------------------------------------------------------------- PDF с текстовым слоем

pdfmetrics.registerFont(TTFont("Arial", FONT_REGULAR))
pdfmetrics.registerFont(TTFont("Arial-Bold", FONT_BOLD))


def text_pdf(path: Path, title: str, lines: list[str], stamp: str | None = None) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setFont("Arial-Bold", 14)
    c.drawCentredString(w / 2, h - 30 * mm, title)
    c.setFont("Arial", 11)
    y = h - 45 * mm
    for line in lines:
        if y < 25 * mm:
            c.showPage()
            c.setFont("Arial", 11)
            y = h - 25 * mm
        c.drawString(20 * mm, y, line)
        y -= 7 * mm
    if stamp:
        c.setStrokeColorRGB(0.1, 0.2, 0.7)
        c.setFillColorRGB(0.1, 0.2, 0.7)
        c.circle(150 * mm, 40 * mm, 18 * mm)
        c.setFont("Arial", 8)
        c.drawCentredString(150 * mm, 40 * mm, stamp)
    c.save()


# ---------------------------------------------------------------- «сканы»

def render_page(title: str, lines: list[str], width: int = 1240, height: int = 1754) -> Image.Image:
    img = Image.new("L", (width, height), 250)
    d = ImageDraw.Draw(img)
    f_title = ImageFont.truetype(FONT_BOLD, 34)
    f_text = ImageFont.truetype(FONT_REGULAR, 26)
    tw = d.textlength(title, font=f_title)
    d.text(((width - tw) / 2, 120), title, font=f_title, fill=20)
    y = 220
    for line in lines:
        d.text((110, y), line, font=f_text, fill=30)
        y += 44
    return img


def scanify(img: Image.Image, seed: int, angle: float) -> Image.Image:
    """Имитация скана: поворот, шум, лёгкое размытие, неравномерная яркость."""
    rnd = random.Random(seed)
    img = img.rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=235)
    px = img.load()
    for _ in range(img.width * img.height // 60):
        x, y = rnd.randrange(img.width), rnd.randrange(img.height)
        px[x, y] = rnd.choice((90, 150, 200))
    img = img.filter(ImageFilter.GaussianBlur(0.7))
    shade = Image.linear_gradient("L").resize(img.size).point(lambda v: 255 - v // 8)
    return Image.composite(img, shade, Image.new("L", img.size, 200))


def image_only_pdf(path: Path, page: Image.Image) -> None:
    """PDF без текстового слоя (скан) — для проверки ветки OCR."""
    buf = io.BytesIO()
    page.convert("RGB").save(buf, format="JPEG", quality=70)
    buf.seek(0)
    c = canvas.Canvas(str(path), pagesize=A4)
    from reportlab.lib.utils import ImageReader

    c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1])
    c.save()


# ---------------------------------------------------------------- пакет

def build(out: Path, today: date) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    edge = out / "edge-cases"
    edge.mkdir(exist_ok=True)
    tax_issued = today - timedelta(days=18)
    tax_valid_until = today + timedelta(days=12)
    app_date = today - timedelta(days=2)
    expected: dict = {"organization": ORG, "inn": INN, "requested_amount": REQUESTED, "today": today.isoformat(), "documents": {}}

    def record(name: str, doc_type: str, fields: dict, note: str = "") -> None:
        expected["documents"][name] = {"type": doc_type, "fields": fields, **({"note": note} if note else {})}

    # 1. Заявление — текстовый PDF
    text_pdf(out / "1.pdf", "ЗАЯВЛЕНИЕ на предоставление субсидии", [
        "В Агентство инноваций Ростовской области",
        f"от {ORG}",
        f"ИНН {INN}   КПП {KPP}   ОГРН {OGRN}",
        f"Юридический адрес: {ADDRESS}",
        "",
        "Прошу предоставить субсидию на реализацию инновационного проекта",
        f"«{PROJECT}»",
        f"в размере {money(REQUESTED)} (восемьсот пятьдесят тысяч рублей 00 копеек).",
        "",
        "Общая стоимость проекта указана в смете расходов (приложение 1).",
        "Достоверность сведений подтверждаю.",
        "",
        f"Генеральный директор  ____________  {DIRECTOR_SHORT}",
        f"Дата: {app_date:%d.%m.%Y}",
    ], stamp=ORG)
    record("1.pdf", "APPLICATION", {"company_name": ORG, "inn": INN, "requested_amount": REQUESTED, "director": DIRECTOR_SHORT, "date": app_date.isoformat()})

    # 2. Выписка ЕГРЮЛ — текстовый PDF, ИНН с ошибкой
    text_pdf(out / "qwe.pdf", "ВЫПИСКА из Единого государственного реестра юридических лиц", [
        f"Дата формирования выписки: {(today - timedelta(days=5)):%d.%m.%Y}",
        "",
        f"Полное наименование: Общество с ограниченной ответственностью «Ромашка-Тех»",
        f"Сокращённое наименование: {ORG}",
        f"ОГРН: {OGRN}",
        f"ИНН: {INN_WRONG}",
        f"КПП: {KPP}",
        f"Адрес юридического лица: {ADDRESS}",
        f"Лицо, имеющее право действовать без доверенности: {DIRECTOR}, генеральный директор",
        "Основной вид деятельности (ОКВЭД): 62.01 Разработка компьютерного программного обеспечения",
        "Уставный капитал: 10 000 руб.",
        "",
        "Выписка сформирована с использованием сервиса ФНС России.",
    ])
    record("qwe.pdf", "REGISTRATION_EXTRACT", {"company_name": ORG, "inn": INN_WRONG, "ogrn": OGRN, "director": DIRECTOR},
           note="ИНН отличается от заявления последней цифрой — должна сработать сверка ИНН между документами")

    # 3. Смета — скан JPG
    est_lines = [f"Организация: {ORG}, ИНН {INN}", f"Проект: {PROJECT}", "", "№   Наименование                         Кол-во      Цена          Сумма"]
    for i, (name, q, p) in enumerate(ESTIMATE_ITEMS, 1):
        est_lines.append(f"{i}.  {name:<36} {q:>5}   {money(p):>14}   {money(q * p):>14}")
    est_lines += ["", f"ИТОГО по смете: {money(ESTIMATE_TOTAL)}", "", f"Генеральный директор  {DIRECTOR_SHORT}"]
    scanify(render_page("СМЕТА РАСХОДОВ ПРОЕКТА", est_lines), seed=4, angle=1.3).convert("RGB").save(out / "scan4.jpg", quality=72)
    record("scan4.jpg", "COST_ESTIMATE", {"total_cost": ESTIMATE_TOTAL, "inn": INN},
           note="Запрошено 850 000, смета 800 000 — разница 50 000 (арифметическая проверка)")

    # 4. Справка ФНС — PDF-скан без текстового слоя, срок истекает через 12 дней
    tax = render_page("СПРАВКА об исполнении обязанности по уплате налогов", [
        "Федеральная налоговая служба",
        f"Выдана: {tax_issued:%d.%m.%Y}",
        f"Налогоплательщик: {ORG}",
        f"ИНН {INN}   КПП {KPP}",
        "",
        "По состоянию на дату выдачи не имеет неисполненной обязанности",
        "по уплате налогов, сборов, страховых взносов, пеней, штрафов.",
        "",
        f"Справка действительна до {tax_valid_until:%d.%m.%Y}.",
    ])
    image_only_pdf(out / "2.pdf", scanify(tax, seed=2, angle=-0.8))
    record("2.pdf", "TAX_CERTIFICATE", {"inn": INN, "date": tax_issued.isoformat(), "valid_until": tax_valid_until.isoformat()},
           note="Скан без текстового слоя (нужен OCR); срок действия истекает через 12 дней — предупреждение")

    # 5. Бизнес-план — DOCX
    doc = Document()
    doc.add_heading(f"Бизнес-план проекта «{PROJECT}»", level=1)
    doc.add_paragraph(f"Заявитель: {ORG}, ИНН {INN}.")
    doc.add_heading("1. Цель проекта", level=2)
    doc.add_paragraph("Сокращение потерь при учёте сырья и готовой продукции за счёт автоматизации.")
    doc.add_heading("2. Финансирование", level=2)
    t = doc.add_table(rows=3, cols=2)
    for row, (k, v) in zip(t.rows, [("Стоимость проекта", money(ESTIMATE_TOTAL)), ("Запрашиваемая субсидия", money(REQUESTED)), ("Собственные средства", money(0))]):
        row.cells[0].text, row.cells[1].text = k, v
    doc.add_heading("3. Срок окупаемости", level=2)
    doc.add_paragraph("18 месяцев.")
    doc.save(out / "doc5.docx")
    record("doc5.docx", "BUSINESS_PLAN", {"total_cost": ESTIMATE_TOTAL, "requested_amount": REQUESTED})

    # 6. Банковские реквизиты — «фото» PNG
    bank = render_page("Карточка предприятия: банковские реквизиты", [
        f"Наименование: {ORG}",
        f"ИНН / КПП: {INN} / {KPP}",
        f"ОГРН: {OGRN}",
        "Расчётный счёт: 40702810952090012345",
        "Банк: ПАО Сбербанк, г. Ростов-на-Дону",
        "БИК: 046015602   Корр. счёт: 30101810600000000602",
    ], height=900)
    scanify(bank, seed=6, angle=2.5).save(out / "IMG_2031.png")
    record("IMG_2031.png", "BANK_DETAILS", {"inn": INN, "account": "40702810952090012345", "bik": "046015602"})

    # 7. Доверенность — текстовый PDF
    text_pdf(out / "7.pdf", "ДОВЕРЕННОСТЬ № 15", [
        f"г. Ростов-на-Дону, {(today - timedelta(days=30)):%d.%m.%Y}",
        "",
        f"{ORG} (ИНН {INN}) в лице генерального директора {DIRECTOR}",
        "уполномочивает Петрову Анну Сергеевну представлять интересы общества",
        "в Агентстве инноваций по вопросам получения субсидии.",
        "",
        f"Доверенность действительна до {(today + timedelta(days=335)):%d.%m.%Y}.",
    ], stamp=ORG)
    record("7.pdf", "POWER_OF_ATTORNEY", {"inn": INN})

    expected["missing"] = ["CONSENT"]
    expected["planted_findings"] = [
        {"code": "MISSING_DOCUMENT", "severity": "ERROR", "document_type": "CONSENT", "message": "Отсутствует согласие на обработку персональных данных"},
        {"code": "AMOUNT_GT_COST", "severity": "ERROR", "message": f"Запрошено {REQUESTED}, смета {ESTIMATE_TOTAL}, разница {REQUESTED - ESTIMATE_TOTAL}"},
        {"code": "INN_MISMATCH", "severity": "WARNING", "message": f"ИНН в заявлении {INN}, в выписке ЕГРЮЛ {INN_WRONG}"},
        {"code": "DOCUMENT_EXPIRES_SOON", "severity": "WARNING", "document_type": "TAX_CERTIFICATE", "message": f"Справка ФНС действительна до {tax_valid_until:%d.%m.%Y} (12 дней)"},
    ]

    # ------------------------------------------------ пограничные файлы
    text_pdf(edge / "Смета расходов (копия).PDF", "СМЕТА РАСХОДОВ ПРОЕКТА", [f"{ORG}, ИНН {INN}", f"ИТОГО по смете: {money(ESTIMATE_TOTAL)}"])
    (edge / "fake.pdf").write_bytes((out / "IMG_2031.png").read_bytes())  # PNG под видом PDF
    (edge / "empty.pdf").write_bytes(b"")
    text_pdf(edge / "zayavlenie.pdf", "СМЕТА РАСХОДОВ ПРОЕКТА", [f"{ORG}, ИНН {INN}", f"ИТОГО по смете: {money(ESTIMATE_TOTAL)}"])
    big = bytearray((out / "1.pdf").read_bytes())
    big += b"%" + b"0" * (int(10.5 * 1024 * 1024) - len(big))  # 10,5 МиБ: выше лимита backend 10 МиБ, ниже лимита nginx
    (edge / "big.pdf").write_bytes(bytes(big))
    (edge / "notes.txt").write_text(f"Смета: итого {ESTIMATE_TOTAL} руб. ИНН {INN}", encoding="utf-8")
    (edge / "virus.exe").write_bytes(b"MZ\x90\x00")
    expected["edge_cases"] = {
        "Смета расходов (копия).PDF": "кириллица и верхний регистр расширения — должен приниматься",
        "fake.pdf": "PNG с расширением .pdf — содержимое не PDF; должен отклоняться или обрабатываться по сигнатуре",
        "empty.pdf": "пустой файл — 400",
        "zayavlenie.pdf": "имя «заявление», внутри смета — тип должен определяться как COST_ESTIMATE",
        "big.pdf": "10,5 МиБ — 413 от backend",
        "notes.txt": "TXT — формат вне MVP (PDF/JPG/PNG), допустим только если поддержан осознанно",
        "virus.exe": "неподдерживаемый формат — 415",
    }

    (out / "expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")
    return expected


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(Path(__file__).parent / "package"))
    parser.add_argument("--today", default=date.today().isoformat(), help="дата «сегодня» для сроков действия")
    args = parser.parse_args()
    result = build(Path(args.out), date.fromisoformat(args.today))
    print(f"Создано {len(result['documents'])} документов пакета и {len(result['edge_cases'])} пограничных файлов в {args.out}")
